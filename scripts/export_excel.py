#!/usr/bin/env python3
"""
大学生助手 v1.1 - Excel 导出引擎
一键导出所有数据到 .xlsx，含图表和格式化

用法：
  python3 export_excel.py                    # 导出到 大学生数据报表.xlsx
  python3 export_excel.py -o 我的报表.xlsx   # 指定文件名
"""

import sys
import os
from datetime import datetime, date, timedelta
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from engine import (
    get_current_semester, get_semester_dir,
    load_json, BASE_DIR, WEEKDAYS_CN,
)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.utils import get_column_letter


# ── 样式 ──
HEADER_FONT = Font(name="微软雅黑", bold=True, size=12, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="667EEA", end_color="667EEA", fill_type="solid")
TITLE_FONT = Font(name="微软雅黑", bold=True, size=16, color="333333")
NORMAL_FONT = Font(name="微软雅黑", size=11)
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)
CENTER = Alignment(horizontal="center", vertical="center")


def style_header(ws, row, cols):
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = THIN_BORDER


def style_data(ws, start_row, end_row, cols):
    for r in range(start_row, end_row + 1):
        for c in range(1, cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = NORMAL_FONT
            cell.border = THIN_BORDER
            cell.alignment = CENTER


def auto_width(ws, cols):
    for col in range(1, cols + 1):
        max_len = 0
        for row in ws.iter_rows(min_col=col, max_col=col):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col)].width = min(max_len * 2 + 4, 40)


def create_excel(output_path: str):
    wb = Workbook()
    sem = get_current_semester()
    sd = get_semester_dir(sem)

    # ── Sheet 1: 课表 ──
    ws = wb.active
    ws.title = "课表"
    ws.merge_cells("A1:E1")
    ws["A1"] = f"📅 课表 - {sem}"
    ws["A1"].font = TITLE_FONT

    headers = ["星期", "时间", "课程", "教师", "地点"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=3, column=i, value=h)
    style_header(ws, 3, len(headers))

    schedule = load_json(sd / "schedule.json")
    for j, c in enumerate(schedule, 4):
        ws.cell(row=j, column=1, value=WEEKDAYS_CN[c["day"] - 1])
        ws.cell(row=j, column=2, value=c["time"])
        ws.cell(row=j, column=3, value=c["course"])
        ws.cell(row=j, column=4, value=c.get("teacher", ""))
        ws.cell(row=j, column=5, value=c.get("location", ""))
    style_data(ws, 4, 3 + len(schedule), len(headers))
    auto_width(ws, len(headers))

    # ── Sheet 2: 作业 ──
    ws = wb.create_sheet("作业")
    ws.merge_cells("A1:F1")
    ws["A1"] = "📝 作业列表"
    ws["A1"].font = TITLE_FONT

    headers = ["ID", "课程", "标题", "截止日期", "状态", "剩余天数"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=3, column=i, value=h)
    style_header(ws, 3, len(headers))

    hw = load_json(sd / "homework.json")
    today = date.today()
    for j, h in enumerate(hw, 4):
        ws.cell(row=j, column=1, value=h["id"])
        ws.cell(row=j, column=2, value=h["course"])
        ws.cell(row=j, column=3, value=h["title"])
        ws.cell(row=j, column=4, value=h["deadline"])
        ws.cell(row=j, column=5, value="✅ 完成" if h["done"] else "⏳ 待完成")
        days = (datetime.strptime(h["deadline"], "%Y-%m-%d").date() - today).days
        ws.cell(row=j, column=6, value=f"{days} 天" if not h["done"] else "-")
    style_data(ws, 4, 3 + len(hw), len(headers))
    auto_width(ws, len(headers))

    # ── Sheet 3: 考试 ──
    ws = wb.create_sheet("考试")
    ws.merge_cells("A1:F1")
    ws["A1"] = "⏰ 考试列表"
    ws["A1"].font = TITLE_FONT

    headers = ["ID", "课程", "类型", "日期", "地点", "倒计时"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=3, column=i, value=h)
    style_header(ws, 3, len(headers))

    exams = load_json(sd / "exam.json")
    for j, e in enumerate(exams, 4):
        ws.cell(row=j, column=1, value=e["id"])
        ws.cell(row=j, column=2, value=e["course"])
        ws.cell(row=j, column=3, value=e["type"])
        ws.cell(row=j, column=4, value=e["date"])
        ws.cell(row=j, column=5, value=e.get("location", ""))
        days = (datetime.strptime(e["date"], "%Y-%m-%d").date() - today).days
        ws.cell(row=j, column=6, value=f"{days} 天" if days >= 0 else "已结束")
    style_data(ws, 4, 3 + len(exams), len(headers))
    auto_width(ws, len(headers))

    # ── Sheet 4: 复习 ──
    ws = wb.create_sheet("复习进度")
    ws.merge_cells("A1:E1")
    ws["A1"] = "📖 知识点复习进度"
    ws["A1"].font = TITLE_FONT

    headers = ["课程", "知识点", "掌握度(%)", "状态", "备注"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=3, column=i, value=h)
    style_header(ws, 3, len(headers))

    review = load_json(sd / "review.json")
    for j, kp in enumerate(review, 4):
        ws.cell(row=j, column=1, value=kp.get("course", ""))
        ws.cell(row=j, column=2, value=kp.get("topic", ""))
        ws.cell(row=j, column=3, value=kp.get("mastery", 0))
        m = kp.get("mastery", 0)
        status = "✅ 掌握" if m >= 80 else ("🟡 熟悉" if m >= 50 else "🔴 薄弱")
        ws.cell(row=j, column=4, value=status)
        ws.cell(row=j, column=5, value=kp.get("notes", ""))
    style_data(ws, 4, 3 + len(review), len(headers))
    auto_width(ws, len(headers))

    # 复习进度饼图
    if review:
        mastered = sum(1 for k in review if k.get("mastery", 0) >= 80)
        weak = sum(1 for k in review if k.get("mastery", 0) < 50)
        mid = len(review) - mastered - weak
        pie = PieChart()
        pie.title = "知识点掌握分布"
        data_ref = Reference(ws, min_col=3, max_col=3, min_row=3, max_row=3 + len(review))
        pie.add_data(data_ref, titles_from_data=True)
        ws.add_chart(pie, "G3")

    # ── Sheet 5: 成长记录 ──
    ws = wb.create_sheet("成长记录")
    ws.merge_cells("A1:G1")
    ws["A1"] = "🌱 每日成长记录"
    ws["A1"].font = TITLE_FONT

    headers = ["日期", "学习时长(h)", "阅读", "运动", "早起", "心情", "日记"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=3, column=i, value=h)
    style_header(ws, 3, len(headers))

    daily = load_json(sd / "daily.json")
    for j, d in enumerate(daily[-30:], 4):  # 最近30天
        ws.cell(row=j, column=1, value=d.get("date", ""))
        ws.cell(row=j, column=2, value=d.get("study_hours", 0))
        ws.cell(row=j, column=3, value="✅" if d.get("reading") else "❌")
        ws.cell(row=j, column=4, value="✅" if d.get("exercise") else "❌")
        ws.cell(row=j, column=5, value="✅" if d.get("early_rise") else "❌")
        ws.cell(row=j, column=6, value=d.get("mood", ""))
        ws.cell(row=j, column=7, value=d.get("diary", ""))
    style_data(ws, 4, 3 + len(daily[-30:]), len(headers))
    auto_width(ws, len(headers))

    # 学习时长趋势图
    if len(daily) >= 2:
        chart = LineChart()
        chart.title = "学习时长趋势"
        chart.y_axis.title = "小时"
        data_ref = Reference(ws, min_col=2, max_col=2, min_row=3, max_row=3 + min(len(daily[-30:]), 30))
        cats = Reference(ws, min_col=1, max_col=1, min_row=4, max_row=3 + min(len(daily[-30:]), 30))
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, "I3")

    # ── Sheet 6: 目标 ──
    ws = wb.create_sheet("目标管理")
    ws.merge_cells("A1:F1")
    ws["A1"] = "🎯 目标管理"
    ws["A1"].font = TITLE_FONT

    headers = ["目标", "类型", "进度(%)", "截止日期", "状态", "里程碑"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=3, column=i, value=h)
    style_header(ws, 3, len(headers))

    goals = load_json(sd / "goals.json")
    for j, g in enumerate(goals, 4):
        ws.cell(row=j, column=1, value=g.get("title", ""))
        ws.cell(row=j, column=2, value={"long": "长期", "short": "短期", "daily": "每日"}.get(g.get("type", ""), ""))
        ws.cell(row=j, column=3, value=g.get("progress", 0))
        ws.cell(row=j, column=4, value=g.get("deadline", ""))
        ws.cell(row=j, column=5, value="🎉 完成" if g.get("progress", 0) >= 100 else "进行中")
        ms = g.get("milestones", [])
        ws.cell(row=j, column=6, value=f"{sum(1 for m in ms if m.get('done'))}/{len(ms)}")
    style_data(ws, 4, 3 + len(goals), len(headers))
    auto_width(ws, len(headers))

    # 目标进度条形图
    if goals:
        chart = BarChart()
        chart.title = "目标完成进度"
        chart.y_axis.title = "%"
        data_ref = Reference(ws, min_col=3, max_col=3, min_row=3, max_row=3 + len(goals))
        cats = Reference(ws, min_col=1, max_col=1, min_row=4, max_row=3 + len(goals))
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, "H3")

    # ── Sheet 7: 刷题统计 ──
    ws = wb.create_sheet("刷题统计")
    ws.merge_cells("A1:F1")
    ws["A1"] = "✏️ 刷题统计"
    ws["A1"].font = TITLE_FONT

    headers = ["日期", "课程", "做题数", "正确数", "正确率(%)", "用时(分钟)"]
    for i, h in enumerate(headers, 1):
        ws.cell(row=3, column=i, value=h)
    style_header(ws, 3, len(headers))

    practice = load_json(sd / "practice.json")
    for j, p in enumerate(practice, 4):
        ws.cell(row=j, column=1, value=p.get("date", ""))
        ws.cell(row=j, column=2, value=p.get("course", ""))
        ws.cell(row=j, column=3, value=p.get("count", 0))
        ws.cell(row=j, column=4, value=p.get("correct", 0))
        rate = round(p.get("correct", 0) / max(p.get("count", 1), 1) * 100, 1)
        ws.cell(row=j, column=5, value=rate)
        ws.cell(row=j, column=6, value=p.get("duration_min", 0))
    style_data(ws, 4, 3 + len(practice), len(headers))
    auto_width(ws, len(headers))

    # 保存
    wb.save(output_path)
    print(f"✅ Excel 报表已生成: {output_path}")
    print(f"📊 包含 {len(wb.sheetnames)} 个 Sheet: {', '.join(wb.sheetnames)}")


def main():
    import argparse
    parser = argparse.ArgumentParser(description="导出 Excel 报表")
    parser.add_argument("-o", "--output", default="大学生数据报表.xlsx", help="输出文件名")
    args = parser.parse_args()
    create_excel(args.output)


if __name__ == "__main__":
    main()
